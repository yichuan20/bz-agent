#!/usr/bin/env python3
from __future__ import annotations
"""
create-pptx.py — generate a PowerPoint presentation from a simple JSON outline.

The agent writes an outline (title + slides), then calls this script to produce
a properly-formatted PPTX file without burning tokens on layout details.

Usage:
  # Outline from a JSON file
  python3 bzcode/scripts/create-pptx.py --outline-file /tmp/outline.json --out /path/to/deck.pptx

  # Inline JSON for small decks
  python3 bzcode/scripts/create-pptx.py --outline '{"title":"Q2 Review","slides":[...]}' --out deck.pptx

Outline JSON schema:
  {
    "title": "Deck Title",
    "author": "Name",          # optional
    "theme": "dark"|"light",   # optional, default "light"
    "slides": [
      {
        "title": "Slide Title",
        "bullets": ["Point 1", "Point 2"],   # optional
        "body": "Free text",                 # optional, used if no bullets
        "notes": "Speaker notes"             # optional
      }
    ]
  }

Output:
  Prints JSON: { "ok": true, "path": "/abs/path/to/deck.pptx", "slides": N }
"""
import argparse
import json
import sys
from pathlib import Path


THEMES = {
    "light": {
        "bg":        (255, 255, 255),
        "title_fg":  (28,  25,  23),
        "body_fg":   (107, 99,  88),
        "accent":    (20, 115, 223),
    },
    "dark": {
        "bg":        (10,  10,  10),
        "title_fg":  (255, 255, 255),
        "body_fg":   (163, 163, 163),
        "accent":    (46, 136, 255),
    },
}


def rgb(triple):
    from pptx.dml.color import RGBColor
    return RGBColor(*triple)


def make_presentation(outline: dict, out_path: Path) -> int:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor

    theme = THEMES.get(outline.get("theme", "light"), THEMES["light"])

    prs = Presentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]  # completely blank

    slides_data = outline.get("slides", [])
    slide_count = 0

    for i, sdata in enumerate(slides_data):
        slide = prs.slides.add_slide(blank_layout)
        slide_count += 1

        # Background
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor(*theme["bg"])

        # Accent bar at top (thin stripe)
        bar = slide.shapes.add_shape(
            1,  # MSO_SHAPE_TYPE rectangle
            Inches(0), Inches(0),
            prs.slide_width, Inches(0.08),
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor(*theme["accent"])
        bar.line.fill.background()

        # Title
        title_text = sdata.get("title", f"Slide {i + 1}")
        tx_title = slide.shapes.add_textbox(
            Inches(0.6), Inches(0.3),
            Inches(12), Inches(1.2),
        )
        tf = tx_title.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        run = p.runs[0]
        run.font.size = Pt(32)
        run.font.bold = True
        run.font.color.rgb = RGBColor(*theme["title_fg"])

        # Slide number (bottom right)
        num_box = slide.shapes.add_textbox(
            Inches(11.8), Inches(6.9),
            Inches(1.4), Inches(0.4),
        )
        ntf = num_box.text_frame.paragraphs[0]
        ntf.text = str(i + 1)
        ntf.alignment = PP_ALIGN.RIGHT
        ntf.runs[0].font.size = Pt(11)
        ntf.runs[0].font.color.rgb = RGBColor(*theme["body_fg"])

        # Body — bullets take priority over body text
        bullets = sdata.get("bullets", [])
        body_text = sdata.get("body", "")

        body_top = Inches(1.7)
        body_box = slide.shapes.add_textbox(
            Inches(0.6), body_top,
            Inches(12), Inches(5.0),
        )
        btf = body_box.text_frame
        btf.word_wrap = True

        if bullets:
            for j, bullet in enumerate(bullets):
                para = btf.paragraphs[0] if j == 0 else btf.add_paragraph()
                para.text = f"• {bullet}"
                para.space_before = Pt(4)
                run = para.runs[0]
                run.font.size = Pt(20)
                run.font.color.rgb = RGBColor(*theme["body_fg"])
        elif body_text:
            btf.paragraphs[0].text = body_text
            run = btf.paragraphs[0].runs[0]
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(*theme["body_fg"])

        # Speaker notes
        notes_text = sdata.get("notes", "")
        if notes_text:
            notes_slide = slide.notes_slide
            notes_slide.notes_text_frame.text = notes_text

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return slide_count


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate a PPTX from a JSON outline.")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--outline",      help="Inline JSON outline string.")
    group.add_argument("--outline-file", help="Path to a JSON outline file.")
    parser.add_argument("--out", required=True, help="Output .pptx path.")
    args = parser.parse_args()

    if args.outline_file:
        f = Path(args.outline_file)
        if not f.exists():
            print(json.dumps({"error": f"outline file not found: {f}"}))
            sys.exit(1)
        try:
            outline = json.loads(f.read_text(encoding="utf-8"))
        except json.JSONDecodeError as e:
            print(json.dumps({"error": f"invalid JSON: {e}"}))
            sys.exit(1)
    else:
        try:
            outline = json.loads(args.outline)
        except json.JSONDecodeError as e:
            print(json.dumps({"error": f"invalid JSON: {e}"}))
            sys.exit(1)

    out_path = Path(args.out)
    if out_path.suffix.lower() != ".pptx":
        out_path = out_path.with_suffix(".pptx")

    try:
        n = make_presentation(outline, out_path)
    except Exception as exc:
        print(json.dumps({"error": f"could not create presentation: {exc}"}))
        sys.exit(1)

    print(json.dumps({"ok": True, "path": str(out_path.resolve()), "slides": n}))


if __name__ == "__main__":
    main()
